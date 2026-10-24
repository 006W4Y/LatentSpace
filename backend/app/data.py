"""Bounded file ingestion, explicit column roles, and local dataset persistence."""
from __future__ import annotations

import csv
import io
import json
import os
import re
import sqlite3
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook
from .presentation import number_series, date_series, column_hints

import numpy as np
import pandas as pd

MAX_ROWS = 20_000
MAX_COLUMNS = 80
MAX_BYTES = 20 * 1024 * 1024
MAX_CELL_CHARS = 1_000_000
csv.field_size_limit(MAX_CELL_CHARS)
ROLES = {"number", "category", "date", "identifier", "text"}


def storage_path() -> Path:
    path = Path(os.environ.get("LATENTSPACE_DATA_DIR", "./data"))
    path.mkdir(parents=True, exist_ok=True)
    return path / "latentspace.sqlite3"


@contextmanager
def connection():
    db = sqlite3.connect(storage_path(), timeout=30)
    db.execute("PRAGMA foreign_keys = ON")
    try:
        with db:
            db.execute("CREATE TABLE IF NOT EXISTS datasets (id TEXT PRIMARY KEY, payload TEXT NOT NULL)")
            yield db
    finally:
        db.close()


def clean_json(value):
    if isinstance(value, dict):
        return {str(k): clean_json(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, np.ndarray)):
        return [clean_json(v) for v in value]
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, (float, np.floating)):
        return float(value) if np.isfinite(value) else None
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.bool_,)):
        return bool(value)
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    return value


def infer_columns(frame: pd.DataFrame, decimal: str) -> list[dict]:
    columns = []
    for name in frame.columns:
        series = frame[name]
        present = series.dropna()
        numbers = number_series(series, decimal)
        numeric = len(present) > 0 and numbers.notna().sum() == len(present)
        name_id = bool(re.search(r"(^id$|[_\s]id$|nummer$|^index$|^uuid$)", name, re.I))
        leading_zero = present.astype(str).str.match(r"^0\d+$").any()
        hints = column_hints(series, decimal)

        if name_id or (numeric and leading_zero):
            kind = "identifier"
        elif len(present) and hints["date_shaped"]:
            kind = "date"
        elif numeric:
            kind = "number"
        elif series.nunique() <= max(30, len(series) * 0.1):
            kind = "category"
        else:
            kind = "text"
        columns.append({"name": name, "role": kind, "missing": int(series.isna().sum()), "unique": int(series.nunique()), "examples": clean_json(present.head(3).tolist()), **hints})
    return columns


def import_file(content: bytes, filename: str, decimal: str = ".") -> dict:
    if not content:
        raise ValueError("Die Datei ist leer.")
    if len(content) > MAX_BYTES:
        raise ValueError("Die Datei ist zu groß. Erlaubt sind maximal 20 MB.")
    suffix = Path(filename).suffix.lower()
    notices = []
    try:
        if suffix == ".csv":
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("cp1252")
                notices.append("Die CSV-Datei wurde mit Windows-Zeichenkodierung gelesen.")
            try:
                delimiter = csv.Sniffer().sniff(text[:16384], delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ";" if decimal == "," else ","
            rows = []
            for row in csv.reader(io.StringIO(text), delimiter=delimiter, strict=True):
                if not row:
                    continue
                if len(row) > MAX_COLUMNS:
                    raise ValueError(f"Bitte höchstens {MAX_COLUMNS} Spalten importieren.")
                if rows and len(row) != len(rows[0]):
                    raise ValueError("Die CSV-Zeilen haben unterschiedliche Spaltenzahlen.")
                if len(rows) > MAX_ROWS:
                    raise ValueError(f"Bitte höchstens {MAX_ROWS:,} Datenzeilen importieren.")
                rows.append(row)
            frame = pd.DataFrame(rows)
        elif suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Die entpackte Excel-Datei überschreitet 100 MB.")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                sheet = book.worksheets[0]
                if (sheet.max_column or 0) > MAX_COLUMNS or (sheet.max_row or 0) > MAX_ROWS + 1:
                    raise ValueError(f"Bitte höchstens {MAX_ROWS:,} Datenzeilen und {MAX_COLUMNS} Spalten im verwendeten Tabellenbereich importieren.")
                sheet.reset_dimensions()  # Do not silently truncate dishonest sheet dimensions.
                rows = []
                for values in sheet.iter_rows(values_only=True):
                    row = list(values)
                    while row and row[-1] is None:
                        row.pop()
                    if len(row) > MAX_COLUMNS:
                        raise ValueError(f"Bitte höchstens {MAX_COLUMNS} Spalten importieren.")
                    if len(rows) > MAX_ROWS:
                        raise ValueError(f"Bitte höchstens {MAX_ROWS:,} Datenzeilen importieren.")
                    if any(isinstance(value, str) and len(value) > MAX_CELL_CHARS for value in row):
                        raise ValueError("Eine Zelle überschreitet eine Million Zeichen.")
                    rows.append(row)
                frame = pd.DataFrame(rows)
                if len(book.sheetnames) > 1:
                    notices.append(f"Importiert wurde das erste Tabellenblatt: {book.sheetnames[0]}. Weitere Blätter sind nicht enthalten.")
            finally:
                book.close()
        else:
            raise ValueError("Bitte eine CSV- oder XLSX-Datei auswählen.")
    except csv.Error as exc:
        raise ValueError("CSV-Struktur oder Zellgröße überschreitet die Importregeln.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Die Datei konnte nicht gelesen werden. Bitte Format und Trennzeichen prüfen.") from exc
    if frame.empty or len(frame) < 2:
        raise ValueError("Die Datei benötigt eine Kopfzeile und mindestens eine Datenzeile.")
    headers = [(str(value).strip() if pd.notna(value) else "") or f"Spalte {i + 1}" for i, value in enumerate(frame.iloc[0])]
    if len(headers) != len(set(headers)):
        raise ValueError("Spaltennamen müssen eindeutig sein. Bitte doppelte Namen in der Kopfzeile ändern.")
    frame = frame.iloc[1:].reset_index(drop=True)
    frame.columns = headers
    frame = frame.map(lambda value: None if isinstance(value, str) and not value.strip() else value)
    validate_size(frame)
    return save_dataset(frame, Path(filename).name, decimal, notices)


def validate_size(frame: pd.DataFrame):
    if len(frame) > MAX_ROWS or len(frame.columns) > MAX_COLUMNS:
        raise ValueError(f"Bitte höchstens {MAX_ROWS:,} Zeilen und {MAX_COLUMNS} Spalten importieren.")
    if frame.empty or not len(frame.columns):
        raise ValueError("Der Datensatz enthält keine Daten.")


def save_dataset(frame: pd.DataFrame, name: str, decimal: str = ".", notices: list | None = None) -> dict:
    validate_size(frame)
    payload = {
        "schema_version": 2, "id": uuid.uuid4().hex, "name": name,
        "created_at": datetime.now(timezone.utc).isoformat(), "decimal": decimal,
        "columns": infer_columns(frame, decimal),
        "rows": clean_json(frame.to_dict("records")), "notices": notices or [],
    }
    with connection() as db:
        db.execute("INSERT INTO datasets VALUES (?, ?)", (payload["id"], json.dumps(payload, allow_nan=False)))
    return metadata(payload)


def metadata(payload: dict) -> dict:
    if any("month_choices" not in column for column in payload["columns"]):
        frame = pd.DataFrame(payload["rows"])
        payload = {**payload, "columns": [{**column, **column_hints(frame[column["name"]], payload["decimal"])} for column in payload["columns"]]}
    return {**{key: value for key, value in payload.items() if key != "rows"}, "row_count": len(payload["rows"]), "preview": payload["rows"][:5]}


def load_dataset(dataset_id: str) -> dict:
    with connection() as db:
        row = db.execute("SELECT payload FROM datasets WHERE id = ?", (dataset_id,)).fetchone()
    if row is None:
        raise KeyError(dataset_id)
    return json.loads(row[0])


def list_datasets() -> list[dict]:
    with connection() as db:
        return [metadata(json.loads(row[0])) for row in db.execute("SELECT payload FROM datasets ORDER BY rowid DESC")]


def delete_dataset(dataset_id: str):
    with connection() as db:
        cursor = db.execute("DELETE FROM datasets WHERE id = ?", (dataset_id,))
    if not cursor.rowcount:
        raise KeyError(dataset_id)


def demo_dataset(kind: str = "retail") -> dict:
    """Preserve the sample entry point; generators live outside file ingestion."""
    from .demos import create_demo
    return create_demo(kind)
